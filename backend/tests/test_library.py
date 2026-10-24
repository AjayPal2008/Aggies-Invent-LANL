"""Training library: institutional / division / role layers, materials, and entry-path rules."""
import io
from pathlib import Path

import pytest

from app.services import library

SAMPLE_PDF = Path(__file__).resolve().parent.parent / "data" / "sample" / "documents" / "project_management_handbook.pdf"


def _path(client, headers, **body):
    r = client.post("/api/library/path", json=body, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_library_is_loaded_in_three_layers(client, admin):
    o = client.get("/api/library", headers=admin).json()
    by_layer = o["counts"]["by_layer"]
    assert by_layer["institutional"] >= 10 and by_layer["role"] >= 30 and by_layer.get("division", 0) >= 2
    assert {p["key"] for p in o["programs"]} == {"weapons", "ste", "operations"}
    assert all(p["requirements"] for p in o["programs"])
    inst = {t["code"]: t for t in o["trainings"] if t["layer"] == "institutional"}
    assert inst["UT-1425"]["frequency"] == "annual" and inst["UT-62676"]["audience"] == "new_hire"
    assert inst["UT-47075"]["refresher_of"] == "UT-9369"
    wrtc = {t["code"]: t for t in o["trainings"] if t["source_system"] == "White Rock Training Center"}
    assert len(wrtc) >= 35 and all(library.allowed_url(m["url"]) for t in wrtc.values() for m in t["materials"] if m["kind"] == "pdf")
    assert sum(m["url"].startswith("https://cdn.lanl.gov/files/") for t in wrtc.values() for m in t["materials"]) >= 38
    roles = {r["key"]: r for r in o["roles"]}
    assert {"project-lead", "engineer", "scientist", "craft-worker"} <= set(roles)
    assert set(roles["project-lead"]["programs"]) == {"weapons", "ste", "operations"}
    assert any(r["kind"] == "soft" for r in roles["project-lead"]["requirements"])


def test_new_hire_gets_all_three_layers(client, manager):
    p = _path(client, manager, program="weapons", role_key="project-lead", entry_path="new_hire")
    assert p["by_layer"]["institutional"] >= 10 and p["by_layer"]["division"] >= 3 and p["by_layer"]["role"] >= 5
    codes = {i["training_code"] for i in p["items"] if i["include"]}
    assert {"UT-62676", "UT-9369", "NFO-1", "NFO-3", "PM-200"} <= codes
    # Refresher courses come on the annual cycle, not in onboarding
    assert "UT-47075" not in {i["training_code"] for i in p["items"]} and "UT-35043" not in codes


def test_internal_move_in_same_division_gets_role_only(client, manager):
    p = _path(client, manager, program="weapons", role_key="project-lead", entry_path="promotion", from_program="weapons")
    assert p["by_layer"] == {"institutional": 0, "division": 0, "role": p["by_layer"]["role"]} and p["by_layer"]["role"] >= 5
    reasons = {i["layer"]: i["why"] for i in p["items"] if not i["include"]}
    assert "annual cycle" in reasons["institutional"] and "Already in this division" in reasons["division"]


def test_move_to_another_division_adds_that_division(client, manager):
    p = _path(client, manager, program="weapons", role_key="project-lead", entry_path="transfer", from_program="ste")
    assert p["by_layer"]["institutional"] == 0 and p["by_layer"]["division"] >= 3 and p["by_layer"]["role"] >= 5


def test_same_role_differs_by_division(client, manager):
    weapons = {i["training_code"] for i in _path(client, manager, program="weapons", role_key="project-lead", entry_path="promotion",
                                                  from_program="weapons")["items"] if i["include"]}
    ste = {i["training_code"] for i in _path(client, manager, program="ste", role_key="project-lead", entry_path="promotion",
                                              from_program="ste")["items"] if i["include"]}
    assert "PM-200" in weapons and "PM-200" in ste  # shared role trainings
    assert "NFO-3" in weapons and "NFO-3" not in ste  # Weapons-only addition


@pytest.mark.parametrize("url,ok", [
    ("https://cdn.lanl.gov/files/its-sm-725-r50_6eb1c.pdf", True),
    ("http://cdn.lanl.gov/files/its-sm-725-r50_6eb1c.pdf", False),
    ("https://evil.example/files/a.pdf", False),
    ("https://cdn.lanl.gov/other/a.pdf", False),
    ("https://cdn.lanl.gov/files/a.pdf?x=1", False),
    ("https://cdn.lanl.gov/files/a.exe", False),
    ("https://www.energy.gov/documents/conduct-operations", True),
    ("https://www.govinfo.gov/content/pkg/CFR-2024-title10-vol4/pdf/CFR-2024-title10-vol4-part835.pdf", True),
    ("http://www.energy.gov/documents/conduct-operations", False),
    ("https://www.energy.gov.evil.example/documents/x", False),
    ("https://docs.nersc.gov/getting-started/?utm=1", False),
])
def test_only_official_public_pdfs_are_fetched(url, ok):
    assert library.allowed_url(url) is ok


def test_import_material_extracts_text_and_dedupes(client, admin, monkeypatch):
    def fake_download(url, target, expect="pdf"):
        target.write_bytes(SAMPLE_PDF.read_bytes())
        return target.stat().st_size

    monkeypatch.setattr(library, "_download", fake_download)
    from app import models
    from app.database import SessionLocal

    with SessionLocal() as db:
        mats = [m for m in db.query(models.TrainingMaterial).filter(models.TrainingMaterial.training_code.in_(["UT-725", "UT-8661"]))]
        for m in mats:
            library.import_material(db, m)
        assert all(m.status == "imported" and m.document_id for m in mats)
        assert len({m.document_id for m in mats}) == 1  # same file twice: one stored copy
        assert db.get(models.Document, mats[0].document_id).pages > 100


def test_upload_slide_deck_and_build_lessons(client, admin):
    from pptx import Presentation

    deck = Presentation()
    for i, (title, body) in enumerate([("Module 1: Ladder Inspection", "Inspect the ladder before each use. Damaged ladders must be tagged out and removed from service."),
                                       ("Module 2: Setting Up", "Extension ladders must extend three feet above the landing. Never stand on the top two rungs."),
                                       ("Module 3: Climbing", "Always maintain three points of contact while climbing. Workers must face the ladder.")]):
        s = deck.slides.add_slide(deck.slide_layouts[1])
        s.shapes.title.text, s.placeholders[1].text = title, body
    buf = io.BytesIO()
    deck.save(buf)
    r = client.post("/api/library/trainings/UT-12985/materials", headers=admin,
                    files={"file": ("ladder.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                    data={"title": "Ladder safety slides"})
    assert r.status_code == 200, r.text
    assert r.json()["kind"] == "pptx" and r.json()["status"] == "uploaded"
    b = client.post("/api/library/trainings/UT-12985/build", headers=admin)
    assert b.status_code == 200, b.text
    t = next(x for x in client.get("/api/library", headers=admin).json()["trainings"] if x["code"] == "UT-12985")
    assert t["modules"] >= 1 and t["layer"] == "role" and t["title"] == "Ladder Safety"  # identity kept


def test_build_needs_material(client, admin):
    assert client.post("/api/library/trainings/UT-42095/build", headers=admin).status_code == 409


def test_only_admins_edit_the_library(client, manager, hire, admin):
    assert client.get("/api/library", headers=manager).status_code == 200
    assert client.get("/api/library", headers=hire).status_code == 403
    assert client.patch("/api/library/trainings/UT-1425", json={"frequency": "once"}, headers=manager).status_code == 403
    r = client.patch("/api/library/trainings/UT-32461", json={"frequency": "annual"}, headers=admin).json()
    assert r["frequency"] == "annual"
    client.patch("/api/library/trainings/UT-32461", json={"frequency": "once"}, headers=admin)
    assert client.patch("/api/library/trainings/UT-1425", json={"layer": "nonsense"}, headers=admin).status_code == 400


def test_edit_division_and_role_requirements(client, admin, manager):
    r = client.post("/api/library/divisions/ste/requirements", json={"training_code": "UT-8876", "phase": 2, "reason": "Cryogens in labs"},
                    headers=admin)
    assert r.status_code == 200
    p = _path(client, manager, program="ste", role_key=None, entry_path="new_hire")
    assert "UT-8876" in {i["training_code"] for i in p["items"] if i["layer"] == "division"}
    assert client.delete(f"/api/library/divisions/requirements/{r.json()['id']}", headers=admin).status_code == 200
    role = client.post("/api/library/roles", json={"title": "Radiation Protection Technician", "family": "Technician",
                                                   "programs": ["operations"]}, headers=admin).json()
    req = client.post(f"/api/library/roles/{role['key']}/requirements", json={"training_code": "UT-47968", "kind": "technical", "phase": 1},
                      headers=admin)
    assert req.status_code == 200
    soft = client.post(f"/api/library/roles/{role['key']}/requirements", json={"training_code": "COM-320", "kind": "soft", "phase": 3},
                       headers=admin)
    assert soft.status_code == 200
    p = _path(client, manager, program="operations", role_key=role["key"], entry_path="promotion", from_program="operations")
    assert {i["training_code"] for i in p["items"] if i["include"]} == {"UT-47968", "COM-320"}


def test_positions_use_library_roles(client, manager):
    roles = {r["key"]: r for r in client.get("/api/library/roles", headers=manager).json()}
    assert any(p["code"] == "CCS-PL3" for p in roles["project-lead"]["positions"])
    plan = client.get("/api/manager/roles/CCS-PL3/plan", headers=manager).json()
    assert any(i["training_code"].startswith("UT-") for i in plan["items"])


def test_annual_cycle_assigns_and_staggers(client, admin):
    from datetime import date, timedelta

    due = (date.today() + timedelta(days=60)).isoformat()
    r = client.post("/api/library/annual-cycle", json={"due_date": due, "stagger_weeks": 4, "codes": ["UT-1425"]}, headers=admin).json()
    assert r["trainings"] == ["UT-1425"] and r["assigned"] + r["reopened"] + r["already_current"] == r["people"]
    again = client.post("/api/library/annual-cycle", json={"due_date": due, "stagger_weeks": 4, "codes": ["UT-1425"]}, headers=admin).json()
    assert again["assigned"] == 0  # running it twice doesn't duplicate


def test_built_lessons_are_drafts_until_published(client, admin, hire):
    from pptx import Presentation

    deck = Presentation()
    for title, body in [("Module 1: Cylinder Storage", "Cylinders must be secured upright with chains or straps at all times."),
                        ("Module 2: Transport", "Valve caps must be in place whenever a cylinder is moved.")]:
        s = deck.slides.add_slide(deck.slide_layouts[1])
        s.shapes.title.text, s.placeholders[1].text = title, body
    buf = io.BytesIO()
    deck.save(buf)
    client.post("/api/library/trainings/UT-9518/materials", headers=admin,
                files={"file": ("cyl.pptx", buf.getvalue(), "application/octet-stream")}, data={"title": "Gas cylinder slides"})
    assert client.post("/api/library/trainings/UT-9518/build", headers=admin).status_code == 200
    # Learners can't see drafts
    hidden = client.get("/api/me/trainings/UT-9518/modules", headers=hire).json()
    assert hidden["modules"] == [] and "reviewed" in hidden["note"]
    review = client.get("/api/library/trainings/UT-9518/modules", headers=admin).json()
    assert review["lessons_status"] == "draft" and review["modules"]
    assert client.post("/api/library/trainings/UT-9518/publish", json={"publish": True}, headers=admin).status_code == 400  # must confirm review
    assert client.post("/api/library/trainings/UT-9518/publish", json={"publish": True, "reviewed": True}, headers=hire).status_code == 403
    assert client.post("/api/library/trainings/UT-9518/publish", json={"publish": True, "reviewed": True}, headers=admin).json()["lessons_status"] == "published"
    assert client.get("/api/me/trainings/UT-9518/modules", headers=hire).json()["modules"]


def _upload_text(client, admin, code, name, text):
    r = client.post(f"/api/library/trainings/{code}/materials", headers=admin,
                    files={"file": (name, text.encode(), "text/plain")}, data={"title": name})
    assert r.status_code == 200, r.text
    return r.json()


def _mats(client, admin, code):
    return next(t for t in client.get("/api/library", headers=admin).json()["trainings"] if t["code"] == code)["materials"]


def test_admin_can_rename_replace_and_delete_an_upload(client, admin, manager):
    m = _upload_text(client, admin, "UT-10922", "old.txt", "Badges must be worn above the waist at all times on LANL property. " * 5)
    assert client.patch(f"/api/library/materials/{m['id']}", json={"title": "Badge rules"}, headers=manager).status_code == 403
    assert client.patch(f"/api/library/materials/{m['id']}", json={"title": "Badge rules"}, headers=admin).json()["title"] == "Badge rules"
    new = client.post(f"/api/library/materials/{m['id']}/file", headers=admin,
                      files={"file": ("new.txt", b"Updated 2026 badge rules. Report lost badges within one hour. " * 5, "text/plain")}).json()
    assert new["id"] == m["id"] and new["document_id"] != m["document_id"] and new["status"] == "uploaded"
    bad = client.post(f"/api/library/materials/{m['id']}/file", headers=admin, files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 400
    assert client.post(f"/api/library/materials/{m['id']}/refresh", headers=admin).status_code == 400  # uploads have no public link
    assert client.delete(f"/api/library/materials/{m['id']}", headers=admin).status_code == 200
    assert all(x["id"] != m["id"] for x in _mats(client, admin, "UT-10922"))


def test_deleting_a_catalog_pdf_hides_it_and_sync_does_not_bring_it_back(client, admin, monkeypatch):
    from app.database import SessionLocal

    mat = _mats(client, admin, "UT-4464")[0]
    assert client.delete(f"/api/library/materials/{mat['id']}", headers=admin).status_code == 200
    assert all(x["id"] != mat["id"] for x in _mats(client, admin, "UT-4464"))
    monkeypatch.setattr(library, "_download", lambda url, target, expect="pdf": (_ for _ in ()).throw(AssertionError("no network in tests")))
    from app import models

    with SessionLocal() as db:
        library.sync_library(db)
        db.commit()
        library.import_all(db, "UT-4464")
        row = db.get(models.TrainingMaterial, mat["id"])
        assert row.status == "removed" and row.document_id is None  # not re-added or re-downloaded
    assert all(x["id"] != mat["id"] for x in _mats(client, admin, "UT-4464"))
    assert client.delete(f"/api/library/materials/{mat['id']}", headers=admin).status_code == 404


def test_redownload_a_linked_pdf(client, admin, monkeypatch):
    def fake_download(url, target, expect="pdf"):
        target.write_bytes(SAMPLE_PDF.read_bytes())
        return target.stat().st_size

    monkeypatch.setattr(library, "_download", fake_download)
    mat = _mats(client, admin, "UT-8876")[0]
    assert client.post(f"/api/library/materials/{mat['id']}/refresh", headers=admin).json()["status"] == "downloading"
    after = next(x for x in _mats(client, admin, "UT-8876") if x["id"] == mat["id"])
    assert after["status"] == "imported" and after["document_id"]


def test_role_detail_gives_job_description_skills_and_lists_per_path(client, manager, admin):
    d = client.get("/api/library/roles/project-lead/detail?program=weapons", headers=manager).json()
    assert d["program"] == "weapons" and d["skills"] and set(d["paths"]) == {"new_hire", "promotion", "transfer"}
    counts = {k: v["included"] for k, v in d["paths"].items()}
    assert counts["new_hire"] > counts["transfer"] > counts["promotion"] > 0
    assert d["paths"]["promotion"]["by_layer"]["institutional"] == 0 and d["paths"]["promotion"]["by_layer"]["division"] == 0
    assert not d["can_edit"]
    assert client.get("/api/library/roles/project-lead/detail?program=nope", headers=manager).status_code == 400
    r = client.patch("/api/library/roles/project-lead", headers=admin, json={
        "title": d["title"], "family": d["family"], "programs": d["programs"], "description": d["description"],
        "job_description": "Leads weapons program projects end to end.", "skills": ["Earned Value Management", "risk management"]})
    assert r.status_code == 200
    d2 = client.get("/api/library/roles/project-lead/detail?program=weapons", headers=manager).json()
    assert d2["job_description"].startswith("Leads") and "earned value management" in {s["skill"] for s in d2["skills"]}


def test_manager_assigns_a_role_list(client, manager):
    people = client.get("/api/manager/team?sort=name", headers=manager).json()
    emp = [p for p in people if p["email"] != "hire@lanl.gov"][-1]  # away from the people other tests use
    d = client.get("/api/library/roles/esh-professional/detail?program=operations", headers=manager).json()
    items = [i for i in d["paths"]["promotion"]["items"] if i["include"]]
    skip = items[0]["training_code"]
    body = {"program": "operations", "entry_path": "promotion", "from_program": "operations", "employee_ids": [emp["id"]], "exclude": [skip]}
    out = client.post("/api/library/roles/esh-professional/assign", json=body, headers=manager).json()
    got = out["people"][0]
    assert skip not in got["assigned"] and set(got["assigned"]) | set(got["already_assigned"]) == {i["training_code"] for i in items[1:]}
    again = client.post("/api/library/roles/esh-professional/assign", json=body, headers=manager).json()["people"][0]
    assert again["assigned"] == []
    assert client.post("/api/library/roles/esh-professional/assign", json={**body, "employee_ids": []}, headers=manager).status_code == 400



def test_every_library_course_has_a_real_source(client, admin):
    """No course is left with only a generic link: each one names a document the AI can read."""
    o = client.get("/api/library", headers=admin).json()
    lib = [t for t in o["trainings"] if t["layer"] in ("institutional", "division", "role")]
    empty = [t["code"] for t in lib if not any(m["kind"] in ("pdf", "web", "pptx", "docx", "text") for m in t["materials"])]
    assert not empty, f"courses without source material: {empty}"
    assert all(library.allowed_url(m["url"]) for t in lib for m in t["materials"] if m["url"] and m["kind"] != "link")


def test_web_pages_become_readable_text(client, admin, monkeypatch):
    from app import models
    from app.database import SessionLocal

    page = "<html><head><script>var x=1</script></head><body><nav>Menu Home</nav><h1>Getting started</h1>" \
           "<p>Never run parallel jobs on a login node. Submit them with sbatch to the batch queue.</p>" * 3 + "</body></html>"

    def fake_download(url, target, expect="pdf"):
        assert expect == "web"
        target.write_text(page)
        return len(page)

    monkeypatch.setattr(library, "_download", fake_download)
    with SessionLocal() as db:
        m = db.query(models.TrainingMaterial).filter_by(training_code="HPC-101", kind="web").first()
        m.document_id, m.status = None, "linked"
        db.commit()
        library.import_material(db, m)
        doc = db.get(models.Document, m.document_id)
        text = " ".join(p.text for p in doc.page_texts)
        assert m.status == "imported" and "login node" in text and "var x" not in text and "Menu Home" not in text


def test_shared_folder_import_places_files_and_respects_deletes(client, admin, tmp_path):
    from app.database import SessionLocal

    (tmp_path / "UT-10922").mkdir()
    (tmp_path / "UT-10922" / "Incident reporting quick guide.txt").write_text(
        "Report every injury to your supervisor before the end of the shift. Call 911 for emergencies. " * 10)
    (tmp_path / "unassigned").mkdir()
    (tmp_path / "unassigned" / "Lab orientation notes.md").write_text("Wear safety glasses in every laboratory at all times. " * 10)
    (tmp_path / "README.md").write_text("not a material")
    with SessionLocal() as db:
        r = library.import_folder(db, tmp_path)
        assert r["imported"] == 2 and r["unassigned"] == 1
        assert library.import_folder(db, tmp_path)["imported"] == 0  # rerun: nothing new
    o = client.get("/api/library", headers=admin).json()
    placed = next(m for t in o["trainings"] if t["code"] == "UT-10922" for m in t["materials"] if m["title"] == "Incident reporting quick guide")
    waiting = next(m for m in o["unassigned"] if m["title"] == "Lab orientation notes")
    assert client.post(f"/api/library/materials/{waiting['id']}/place", json={"training_code": "UT-32461"}, headers=admin).status_code == 200
    assert all(m["id"] != waiting["id"] for m in client.get("/api/library", headers=admin).json()["unassigned"])
    # The admin's own files come first when lessons are built
    with SessionLocal() as db:
        assert library.primary_document(db, "UT-10922").title == "Incident reporting quick guide"
    client.delete(f"/api/library/materials/{placed['id']}", headers=admin)
    with SessionLocal() as db:
        library.import_folder(db, tmp_path)  # deleted catalog/folder files don't come back
    assert all(m["title"] != "Incident reporting quick guide" for t in client.get("/api/library", headers=admin).json()["trainings"]
               if t["code"] == "UT-10922" for m in t["materials"])


def test_export_copies_uploads_into_the_shared_folder(client, admin, tmp_path):
    from app.database import SessionLocal

    _upload_text(client, admin, "UT-57950", "cui-notes.txt", "Mark CUI documents with the CUI banner at the top and bottom of each page. " * 5)
    with SessionLocal() as db:
        r = library.export_to_folder(db, tmp_path)
        assert r["copied"] >= 1 and any(f.startswith("UT-57950/") for f in r["files"])
        assert library.export_to_folder(db, tmp_path)["copied"] == 0  # already there


def test_build_missing_and_publish_all(client, admin, hire):
    from app.database import SessionLocal

    _upload_text(client, admin, "UT-32461", "ems-basics.txt",
                 "Module 1: Aspects. Workers must identify environmental aspects of their work before starting. " * 6)
    with SessionLocal() as db:
        out = library.build_missing(db, codes=["UT-32461"])
        assert out["built"] == ["UT-32461"]
    o = client.get("/api/library", headers=admin).json()
    assert o["lessons"]["draft"] >= 1
    assert client.post("/api/library/publish-all", json={"codes": ["UT-32461"]}, headers=admin).status_code == 400
    assert client.post("/api/library/publish-all", json={"codes": ["UT-32461"], "reviewed": True}, headers=hire).status_code == 403
    r = client.post("/api/library/publish-all", json={"codes": ["UT-32461"], "reviewed": True}, headers=admin).json()
    assert r["published"] == ["UT-32461"]


def test_unbuilt_lessons_still_use_the_library_text(client, admin):
    from app import models
    from app.database import SessionLocal
    from app.services import ai

    _upload_text(client, admin, "UT-21208", "export-basics.txt", "An export includes releasing technology to a foreign person in the United States. " * 5)
    library._EXCERPTS.clear()
    with SessionLocal() as db:
        t = db.query(models.Training).filter_by(code="UT-21208").one()
        block = ai._training_block(t)
    assert "<source>" in block and "foreign person" in block
